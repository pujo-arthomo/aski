#!/usr/bin/env python3
"""
extract_items.py — generate supabase/seed_items.sql dari file xlsx form ASKI asli.

Kenapa script ini ada: item_pertanyaan & pilihan_jawaban isinya ratusan baris
pertanyaan resmi ASKI (teks pernyataan, 5 pilihan a-e per item). Daripada
diketik ulang manual (rawan typo, dan harus diulang tiap instrumen berubah),
script ini membaca langsung dari sheet kerja form ASKI ("ASKI UP1" / "ASKI UK")
dan menghasilkan SQL INSERT siap jalan.

Cara pakai:
    pip install openpyxl
    python3 scripts/extract_items.py \
        --file "1. FORM ASKI UP1 2026.xlsx" --sheet "ASKI UP1" --jenis UP \
        --file "3. FORM ASKI UK 2026.xlsx"  --sheet "ASKI UK"  --jenis UK \
        --out supabase/seed_items.sql

Lalu jalankan supabase/schema.sql dulu (isi aspek_penilaian & sub_aspek),
baru jalankan supabase/seed_items.sql setelahnya di SQL Editor Supabase.

Cara kerja singkat:
- Form ASKI menandai sub-aspek dengan baris "1.1. SUB-ASPEK PENCIPTAAN" dst,
  dan bagian dengan baris "A. BAGIAN ELEKTRONIK" / "B. BAGIAN KONVENSIONAL".
- Tiap item pertanyaan adalah baris bernomor ("1.", "2.", dst di kolom A) yang
  kolom levelnya (kolom F) sudah terisi angka — itu tandanya baris "induk",
  bukan baris pilihan.
- Baris pilihan jawabannya sendiri ada di bawahnya, bertanda "a." s.d. "e."
  di kolom A.
- Skor tiap level SELALU mengikuti tabel referensi yang sama di semua form
  ASKI: level 0->skor 0, level 1->20, level 2->50, level 3->70, level 4->100,
  dan urutan pilihan SELALU a=level 0 ... e=level 4. Ini dicek konsisten di
  puluhan contoh pada form UP1 & UK 2026, jadi dipakai sebagai aturan umum.
- Item yang jumlah pilihannya bukan 5 (jarang terjadi) DILEWATI dan dicatat
  di bagian akhir sebagai "PERLU DICEK MANUAL" — supaya tidak ada data yang
  diam-diam salah masuk.
"""

import argparse
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Perlu openpyxl. Jalankan: pip install openpyxl")

LEVEL_SKOR = {0: 0, 1: 20, 2: 50, 3: 70, 4: 100}
KODE_URUT = ["a", "b", "c", "d", "e"]

SUBASPEK_NUM_RE = re.compile(r"^\d+\.\d+\.?$")
ITEM_NO_RE = re.compile(r"^\d+\.$")
OPT_KODE_RE = re.compile(r"^([a-e])\.$")


def sql_escape(s):
    return (s or "").strip().replace("'", "''")


def normalize_subaspek_name(raw):
    """Samakan penamaan sub-aspek dari sheet form dengan nama di schema.sql."""
    raw = raw.strip().title()
    fixes = {
        "Sumber Daya Manusia Kearsipan": "Sumber Daya Manusia Kearsipan",
        "Sarana Dan Prasarana": "Sarana dan Prasarana Kearsipan",
        "Sarana Dan Prasarana Kearsipan": "Sarana dan Prasarana Kearsipan",
        "Pengendalian Naskah Dinas": "Pengendalian Naskah Dinas",
        "Penggunaan Arsip": "Penggunaan",
        "Penggunaan": "Penggunaan",
        "Penciptaan": "Penciptaan",
        "Pemeliharaan": "Pemeliharaan",
        "Penyusutan": "Penyusutan",
    }
    return fixes.get(raw, raw)


def build_merge_lookup(ws):
    """Beberapa sheet form ASKI pakai merged cell untuk kolom LEVEL/SKOR — nilainya
    cuma ada di sel kiri-atas rentang merge, sel lain di rentang itu None kalau
    dibaca langsung. Ini bikin lookup supaya nilai itu tetap kebaca di baris
    manapun dalam rentang merge-nya."""
    lookup = {}
    for rng in ws.merged_cells.ranges:
        top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                lookup[(r, c)] = top_left
    return lookup


def extract_sheet(path, sheet_name, jenis):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    merged = build_merge_lookup(ws)

    def cell_val(row_num, col_idx, direct):
        """direct = nilai yang terbaca langsung dari cell.value. Kalau None dan
        cell ini bagian dari merge range, pakai nilai merge-nya."""
        if direct is not None:
            return direct
        return merged.get((row_num, col_idx))

    # Pra-scan: kadang LEVEL/SKOR sebuah item "nyasar" ke baris "A./B. BAGIAN ..."
    # tepat di atasnya, bukan di baris pernyataan itemnya sendiri (kelihatannya
    # salah format manual di template aslinya, bukan pola yang konsisten).
    # Simpan semua baris yang punya angka di kolom LEVEL+SKOR supaya bisa dipakai
    # sebagai fallback kalau baris item sendiri kosong.
    numeric_fg_by_row = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        rn = row[0].row
        fv = row[5].value if len(row) > 5 else None
        gv = row[6].value if len(row) > 6 else None
        if isinstance(fv, (int, float)) and isinstance(gv, (int, float)):
            numeric_fg_by_row[rn] = (fv, gv)

    items = []
    skipped = []
    current_sub = None
    current_bagian = None
    sub_item_counter = {}

    pending = None  # item lagi dikumpulkan pilihannya

    def flush_pending():
        if pending is None:
            return
        opts = pending["options"]
        if not pending.get("level_ok", True):
            skipped.append((pending["nomor"], pending["pernyataan"][:70], "level bukan angka standar (skema skor khusus)"))
            return
        if len(opts) != 5 or [o[0] for o in opts] != KODE_URUT:
            skipped.append((pending["nomor"], pending["pernyataan"][:70], f"{len(opts)} pilihan (bukan 5 a-e)"))
            return
        items.append(pending)

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        row_num = row[0].row
        a = row[0].value
        b = row[1].value if len(row) > 1 else None
        f_direct = row[5].value if len(row) > 5 else None
        f = cell_val(row_num, 6, f_direct)

        if isinstance(a, str) and SUBASPEK_NUM_RE.match(a.strip()) and isinstance(b, str) and b.strip().upper().startswith("SUB-ASPEK"):
            flush_pending()
            pending = None
            current_sub = normalize_subaspek_name(b.strip()[len("SUB-ASPEK"):].strip())
            current_bagian = None
            sub_item_counter[current_sub] = 0
            continue

        if isinstance(a, str) and a.strip() in ("A.", "B.") and isinstance(b, str) and "BAGIAN" in b.upper():
            flush_pending()
            pending = None
            current_bagian = "elektronik" if a.strip() == "A." else "konvensional"
            continue

        # baris induk item: "1." (kadang kesimpan sebagai angka murni 1.0, bukan
        # teks "1.") + teks pernyataan + level numerik di kolom F
        is_item_no = (isinstance(a, str) and ITEM_NO_RE.match(a.strip())) or (
            isinstance(a, (int, float)) and float(a).is_integer()
        )
        if is_item_no and isinstance(b, str) and len(b) > 12 and current_sub is not None:
            if not isinstance(f, (int, float)):
                for back in (1, 2, 3):
                    if (row_num - back) in numeric_fg_by_row:
                        f = numeric_fg_by_row[row_num - back][0]  # fallback: nyasar ke baris di atasnya
                        break
            # Selalu tutup item sebelumnya & mulai item baru di sini — kalau tidak,
            # baris pilihan a-e milik item ini bisa "nyasar" nempel ke pending
            # lama saat level-nya tidak bisa dipastikan angka (lihat komentar level_ok).
            flush_pending()
            sub_item_counter[current_sub] = sub_item_counter.get(current_sub, 0) + 1
            no_local = sub_item_counter[current_sub]
            bagian_code = {"elektronik": "A", "konvensional": "B"}.get(current_bagian, "")
            nomor = f"{current_sub[:3]}-{bagian_code}.{no_local}" if bagian_code else f"{current_sub[:3]}.{no_local}"
            pending = {
                "sub_aspek": current_sub,
                "bagian": current_bagian,
                "nomor": nomor,
                "pernyataan": b.strip(),
                "options": [],
                "level_ok": isinstance(f, (int, float)),
            }
            continue

        # baris pilihan: "a." s.d. "e."
        if isinstance(a, str) and pending is not None:
            m = OPT_KODE_RE.match(a.strip())
            if m and isinstance(b, str) and b.strip():
                pending["options"].append((m.group(1), b.strip()))
                continue

    flush_pending()

    return items, skipped


def to_sql(all_items):
    lines = [
        "-- ============================================================================",
        "-- SEED: ITEM_PERTANYAAN & PILIHAN_JAWABAN",
        "-- Auto-generated oleh scripts/extract_items.py — JANGAN edit manual di sini,",
        "-- edit sumbernya (file xlsx form ASKI) lalu generate ulang.",
        "-- Jalankan file ini SETELAH supabase/schema.sql (butuh aspek_penilaian & sub_aspek).",
        "-- ============================================================================",
        "",
    ]
    for idx, it in enumerate(all_items, start=1):
        bagian_sql = f"'{it['bagian']}'" if it["bagian"] else "null"
        lines.append(f"-- {idx}. [{it['jenis']}] {it['sub_aspek']} — {it['nomor']}")
        lines.append("with itm as (")
        lines.append("  insert into item_pertanyaan (sub_aspek_id, nomor, bagian, pernyataan, urutan)")
        lines.append("  select sa.id, "
                      f"'{sql_escape(it['nomor'])}', {bagian_sql}, "
                      f"'{sql_escape(it['pernyataan'])}', {idx}")
        lines.append("  from sub_aspek sa join aspek_penilaian ap on ap.id = sa.aspek_id")
        lines.append(f"  where ap.jenis_instrumen = '{it['jenis']}' and sa.nama = '{sql_escape(it['sub_aspek'])}'")
        lines.append("  returning id")
        lines.append(")")
        values = []
        for kode, label in it["options"]:
            level = KODE_URUT.index(kode)
            skor = LEVEL_SKOR[level]
            values.append(f"select id, '{kode}', '{sql_escape(label)}', {level}, {skor} from itm")
        lines.append("insert into pilihan_jawaban (item_id, kode, label, level, skor)")
        lines.append("\n  union all ".join(values) + ";")
        lines.append("")
    return "\n".join(lines)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--file", action="append", required=True, help="path xlsx (bisa diulang)")
    p.add_argument("--sheet", action="append", required=True, help="nama sheet, urutan sejajar dengan --file")
    p.add_argument("--jenis", action="append", required=True, choices=["UP", "UK"], help="urutan sejajar dengan --file")
    p.add_argument("--out", default="supabase/seed_items.sql")
    args = p.parse_args()

    if not (len(args.file) == len(args.sheet) == len(args.jenis)):
        sys.exit("--file, --sheet, --jenis harus jumlahnya sama dan berurutan.")

    all_items = []
    all_skipped = []
    for path, sheet, jenis in zip(args.file, args.sheet, args.jenis):
        items, skipped = extract_sheet(path, sheet, jenis)
        for it in items:
            it["jenis"] = jenis
        all_items.extend(items)
        all_skipped.extend([(jenis, *s) for s in skipped])
        print(f"[{jenis}] {sheet}: {len(items)} item terambil, {len(skipped)} dilewati")

    sql = to_sql(all_items)
    with open(args.out, "w", encoding="utf-8") as f:
        f.write(sql)
    print(f"\nSelesai -> {args.out} ({len(all_items)} item total)")

    if all_skipped:
        print("\nPERLU DICEK MANUAL (bukan 5 pilihan a-e, kemungkinan tabel kriteria/kalkulator khusus):")
        for jenis, nomor, teks, alasan in all_skipped:
            print(f"  [{jenis}] {nomor} — {teks}... ({alasan})")


if __name__ == "__main__":
    main()
